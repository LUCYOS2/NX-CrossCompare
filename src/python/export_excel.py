"""C++ 코어가 만든 비교 결과 CSV를 서식 있는 Excel(.xlsx)로 변환한다.

역할 분리(개발계획_v2.md §2): 계산은 C++(Rule Engine), 엑셀 서식/리포트는 Python.

CSV 컬럼(ComparisonExporter.cpp 참고, 세로형/long - 규칙×인치 조합마다 한 행):
    rule, model, inch, ctq_code, check_point_category, spec_mm, tolerance_plus_mm,
    tolerance_minus_mm, measurement_type, image_path, within_tolerance

목표(2026-09-12 요청): "치수 비교 테이블의 규칙관리 측정부위, 이미지 등 정보와 측정된
데이터를 같이 추출". 화면(치수 비교 테이블)은 규칙×인치 매트릭스를 그대로 쓰지만, 엑셀은
사용자가 지정한 8개 컬럼(규칙/모델/인치/포인트(부위)/관리항목 종류/도면 스펙/공차/
측정방법)의 세로형 표로 뽑는다 - 여기서 "도면 스펙"은 목표(nominal)값이 아니라 실측값
(value_mm/spec_mm)을 그대로 보여주는 것으로 확정됐다(사용자 확인 완료).
추가로 각 규칙의 측정 부위 캡쳐 이미지(image_path)를 맨 앞 컬럼에 삽입한다.

사용법:
    python export_excel.py comparison_export.csv comparison_export.xlsx

주의: 개인PC에 Python/openpyxl이 설치되어 있지 않아 아직 실행 검증을 하지 못한 초안입니다.
회사PC(또는 Python 설치 후)에서 처음 실행할 때 한 번 동작 확인이 필요합니다.
"""

import csv
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "openpyxl이 필요합니다. 'pip install openpyxl' 실행 후 다시 시도하세요."
    ) from exc

PASS_FILL = PatternFill(start_color="C8FFC8", end_color="C8FFC8", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC8C8", end_color="FFC8C8", fill_type="solid")
HEADER_FONT = Font(bold=True)

# 사용자가 지정한 8개 컬럼(규칙/모델/인치/포인트(부위)/관리항목 종류/도면 스펙/공차/
# 측정방법) 앞에 이미지 컬럼을 하나 더 둔다 - "이미지 등 정보"도 같이 뽑히길 원했으므로.
HEADERS = [
    "측정 부위 이미지", "규칙", "모델", "인치", "포인트(부위)", "관리항목 종류",
    "도면 스펙(mm)", "공차(mm)", "측정방법",
]
IMAGE_COL_WIDTH_CHARS = 20
IMAGE_ROW_HEIGHT_POINTS = 90
IMAGE_MAX_WIDTH_PX = 160


def _format_tolerance(plus_mm: str, minus_mm: str) -> str:
    plus = float(plus_mm)
    minus = float(minus_mm)
    if plus == minus:
        return f"±{plus:g}"
    return f"+{plus:g}/-{minus:g}"


def convert(csv_path: str, xlsx_path: str) -> None:
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    wb = Workbook()
    ws = wb.active
    ws.title = "치수 비교"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT

    image_col_letter = get_column_letter(1)
    ws.column_dimensions[image_col_letter].width = IMAGE_COL_WIDTH_CHARS

    csv_dir = os.path.dirname(os.path.abspath(csv_path))

    for row in rows:
        within_tolerance = row["within_tolerance"] == "TRUE"
        row_idx = ws.max_row + 1
        # § 공용부품(ALL) 표시(2026-09-12) - C++ 쪽(ImportInchDialog::kAllInchValue)이
        # "모든 인치에 공통" 부품을 1000이라는 특수 인치값으로 넘긴다. 숫자 1000을 그대로
        # 보여주면 혼동되므로 "ALL" 문자열로 바꿔 보여준다.
        inch_value = int(row["inch"])
        inch_display = "ALL" if inch_value == 1000 else inch_value
        ws.append([
            "",  # 이미지는 아래에서 add_image로 따로 삽입
            row["rule"],
            row["model"],
            inch_display,
            row["ctq_code"],
            row["check_point_category"],
            float(row["spec_mm"]),
            _format_tolerance(row["tolerance_plus_mm"], row["tolerance_minus_mm"]),
            row["measurement_type"],
        ])
        ws.row_dimensions[row_idx].height = IMAGE_ROW_HEIGHT_POINTS

        # "도면 스펙" 칸에 공차 통과/실패 색상을 입혀서 한눈에 보이게.
        spec_col = HEADERS.index("도면 스펙(mm)") + 1
        ws.cell(row=row_idx, column=spec_col).fill = PASS_FILL if within_tolerance else FAIL_FILL

        # § 측정 부위 이미지 삽입 - RuleEditorDialog에서 캡쳐한 jpg를 그대로 셀에 넣는다.
        # 경로가 비어있거나(이미지 없음) 파일이 실제로 없으면(다른 PC에서 만든 경로 등)
        # 조용히 스킵 - 엑셀 생성 자체가 깨지면 안 되므로.
        image_path = row["image_path"]
        if image_path:
            resolved_path = image_path if os.path.isabs(image_path) else os.path.join(csv_dir, image_path)
            if os.path.isfile(resolved_path):
                try:
                    img = XlImage(resolved_path)
                    if img.width > IMAGE_MAX_WIDTH_PX:
                        scale = IMAGE_MAX_WIDTH_PX / img.width
                        img.width = IMAGE_MAX_WIDTH_PX
                        img.height = int(img.height * scale)
                    ws.add_image(img, f"{image_col_letter}{row_idx}")
                except Exception as exc:  # 손상된 이미지 등 - 행 전체를 막지 않는다.
                    print(f"경고: 이미지 삽입 실패({image_path}): {exc}")

    for col_idx, header in enumerate(HEADERS, start=1):
        if col_idx == 1:
            continue  # 이미지 열은 위에서 이미 폭을 정함
        letter = get_column_letter(col_idx)
        max_length = max(
            [len(str(header))]
            + [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)]
        )
        ws.column_dimensions[letter].width = max_length + 2

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    wb.save(xlsx_path)
    print(f"저장 완료: {xlsx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("사용법: python export_excel.py <입력.csv> <출력.xlsx>")
    convert(sys.argv[1], sys.argv[2])
